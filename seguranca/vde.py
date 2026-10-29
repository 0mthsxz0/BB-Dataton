"""Adaptador do VDE anual: preserva unidade, abrangência e falta de município.

Registros estaduais e vítimas ficam consultáveis na tabela vde_registros,
mas nunca são distribuídos entre municípios ou tratados como ocorrências.
"""
from collections import Counter
from datetime import datetime, timezone
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .banco import crime_canonico, normalizar, numero, registrar_fonte, sigla_uf


COLUNAS = ['uf', 'municipio', 'evento', 'data_referencia', 'agente', 'arma',
           'faixa_etaria', 'feminino', 'masculino', 'nao_informado',
           'total_vitima', 'total', 'total_peso', 'abrangencia']


def ler_vde(arquivo):
    livro = load_workbook(arquivo, read_only=True, data_only=True)
    try:
        if len(livro.sheetnames) != 1:
            raise ValueError('VDE anual exige uma aba; revisar o novo layout.')
        folha = livro.active
        linhas = folha.iter_rows(values_only=True)
        cabecalho = list(next(linhas))
        if cabecalho != COLUNAS:
            raise ValueError(f'Colunas VDE inesperadas: {cabecalho}')
        categorias, abrangencias, meses = Counter(), Counter(), set()
        roubos, total = [], 0
        for linha, valores in enumerate(linhas, 2):
            if not any(v is not None for v in valores):
                continue
            total += 1
            r = dict(zip(cabecalho, valores))
            categorias[str(r['evento'])] += 1
            abrangencias[str(r['abrangencia'])] += 1
            data = pd.Timestamp(r['data_referencia'])
            if pd.isna(data) or data.day != 1:
                raise ValueError(f'Data VDE inválida na linha {linha}')
            meses.add(data.strftime('%Y-%m'))
            if crime_canonico(r['evento']).startswith('roubo '):
                r['linha'] = linha
                r['data_referencia'] = data.strftime('%Y-%m-%d')
                roubos.append(r)
        if len({m[:4] for m in meses}) != 1:
            raise ValueError('VDE anual vazio ou com mais de um ano; revisar o arquivo.')
        auditoria = {'linhas_arquivo': total, 'colunas': cabecalho,
                     'categorias': dict(sorted(categorias.items())),
                     'abrangencias': dict(sorted(abrangencias.items())),
                     'meses': sorted(meses)}
        return roubos, auditoria
    finally:
        livro.close()


def gravar_vde(con, arquivo, url, registros, auditoria, somente_bruto=False):
    """Importação atômica, idempotente e sem soma de dimensões desconhecidas."""
    if not somente_bruto and not con.execute('SELECT 1 FROM municipios LIMIT 1').fetchone():
        raise ValueError('Importe o cadastro IBGE antes do VDE.')
    con.execute('''CREATE TABLE IF NOT EXISTS vde_registros (
        fonte_id INTEGER NOT NULL REFERENCES fontes(id), linha INTEGER NOT NULL,
        uf TEXT NOT NULL, municipio_original TEXT, codigo_ibge TEXT REFERENCES municipios(codigo_ibge),
        crime TEXT NOT NULL, mes TEXT NOT NULL, abrangencia TEXT NOT NULL,
        metrica TEXT NOT NULL, valor INTEGER, status TEXT NOT NULL, dados TEXT NOT NULL,
        PRIMARY KEY(fonte_id,linha))''')
    fonte = registrar_fonte(con, arquivo, url)
    if somente_bruto and con.execute('SELECT 1 FROM ocorrencia_fontes WHERE fonte_id=? LIMIT 1', (fonte,)).fetchone():
        raise ValueError('Fonte já publicada como municipal; coleta bruta não pode rebaixar sua auditoria.')
    por_nome = {}
    for r in con.execute('SELECT codigo_ibge,municipio,uf FROM municipios'):
        por_nome.setdefault((r['uf'], normalizar(r['municipio'])), []).append(r['codigo_ibge'])
    chaves = Counter((sigla_uf(r['uf']), normalizar(r['municipio']),
                      crime_canonico(r['evento']), r['data_referencia'])
                     for r in registros if r['abrangencia'] == 'Estadual')
    resumo = Counter()
    perfis = {}
    preparadas = []
    por_uf = {}
    for r in registros:
        uf, crime, mes = sigla_uf(r['uf']), crime_canonico(r['evento']), r['data_referencia']
        if mes >= datetime.now(timezone.utc).strftime('%Y-%m-01'):
            raise ValueError('VDE contém mês corrente ou futuro.')
        candidatos = por_nome.get((uf, normalizar(r['municipio'])), [])
        codigo = candidatos[0] if len(candidatos) == 1 and not somente_bruto else None
        n, vitimas = numero(r['total'], inteiro=True), numero(r['total_vitima'], inteiro=True)
        if n is not None and vitimas is not None:
            raise ValueError('VDE com duas unidades na mesma linha; revisar o adaptador.')
        metrica = 'ocorrencias' if n is not None else 'vitimas' if vitimas is not None else 'ausente'
        valor = n if n is not None else vitimas
        repetida = chaves[(uf, normalizar(r['municipio']), crime, mes)] > 1
        if r['abrangencia'] != 'Estadual':
            status = 'abrangencia_nao_estadual'
        elif normalizar(r['municipio']) in ('nao informado', ''):
            status = 'sem_municipio'
        elif somente_bruto:
            status = 'aguardando_validacao_municipal'
        elif codigo is None:
            status = 'municipio_sem_correspondencia'
        elif repetida:
            status = 'chave_municipal_repetida'
        elif any(r[c] is not None and str(r[c]).strip() for c in ('agente', 'arma', 'faixa_etaria')):
            status = 'dimensao_nao_agregada'
        elif metrica != 'ocorrencias':
            status = 'vitimas_nao_sao_ocorrencias' if metrica == 'vitimas' else 'valor_ausente'
        else:
            status = 'ocorrencia_municipal'
            preparadas.append(((codigo, crime, mes), n))
        resumo[status] += 1
        perfil = perfis.setdefault(r['evento'], {'linhas': 0, 'com_municipio': 0,
                   'sem_municipio': 0, 'ocorrencias_preenchidas': 0, 'vitimas_preenchidas': 0})
        perfil['linhas'] += 1
        perfil['com_municipio'] += int(codigo is not None)
        perfil['sem_municipio'] += int(codigo is None)
        perfil['ocorrencias_preenchidas'] += int(n is not None)
        perfil['vitimas_preenchidas'] += int(vitimas is not None)
        uf_status = por_uf.setdefault(uf, Counter())
        uf_status[status] += 1
        con.execute('INSERT OR REPLACE INTO vde_registros VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',
                    (fonte, r['linha'], uf, r['municipio'], codigo, crime, mes,
                     r['abrangencia'], metrica, valor, status, json.dumps(r, ensure_ascii=False, default=str)))
    for chave, n in preparadas:
        existente = con.execute('SELECT ocorrencias FROM ocorrencias WHERE codigo_ibge=? AND crime=? AND mes=?', chave).fetchone()
        if existente and existente[0] != n:
            raise RuntimeError(f'Fontes conflitantes para {chave}; revisão manual necessária.')
        con.execute('INSERT OR IGNORE INTO ocorrencias VALUES(?,?,?,?)', (*chave, n))
        con.execute('INSERT OR IGNORE INTO ocorrencia_fontes VALUES(?,?,?,?)', (*chave, fonte))
    con.execute('DELETE FROM indices')
    resultado = {**auditoria, 'fonte_id': fonte, 'linhas_roubo': len(registros),
                 'status': dict(resumo), 'por_crime': perfis,
                 'por_uf': {uf: dict(c) for uf, c in sorted(por_uf.items())},
                 'observacoes_municipais_importadas': len(preparadas),
                 'chaves_estaduais_repetidas': sum(n > 1 for n in chaves.values()),
                 'linhas_estaduais_em_chaves_repetidas': sum(n for n in chaves.values() if n > 1),
                 'limite': 'Sem rateio estadual, sem somar vítimas a ocorrências e sem preencher ausências com zero.'}
    con.execute('UPDATE fontes SET metadados=? WHERE id=?',
                (json.dumps(resultado, ensure_ascii=False), fonte))
    return resultado


def importar_vde(con, arquivo, url, somente_bruto=False):
    registros, auditoria = ler_vde(Path(arquivo))
    return gravar_vde(con, arquivo, url, registros, auditoria, somente_bruto)
